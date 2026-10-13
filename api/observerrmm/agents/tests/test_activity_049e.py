"""049-E · monitoreo en vivo (captura forzada), URL y marca manual en la ingesta,
exportación XLSX, horario laboral y canon de procs_i."""

import tempfile
from io import BytesIO
from unittest.mock import AsyncMock, patch

from django.core.cache import cache
from django.test import override_settings
from django.utils import timezone as djangotime
from model_bakery import baker

from agents.activity import canon_activity_day, parse_work_days, parse_work_hours
from agents.models import AgentActivityDaily, AgentActivityScreenshot
from core.utils import get_core_settings
from logs.models import AuditLog
from observerrmm.test import ObserverTestCase


class TestHelpers049E(ObserverTestCase):
    def test_horario_laboral(self):
        self.assertEqual(parse_work_hours("08:00-18:00"), (480, 1080))
        self.assertIsNone(parse_work_hours("18:00-08:00"))
        self.assertIsNone(parse_work_hours("ocho a seis"))
        self.assertEqual(parse_work_days("0, 1,2,9,x"), {0, 1, 2})

    def test_canon_de_procs_i(self):
        dia = {"usuarios": [{"procs_i": {"chrome.exe": 3, "chrome": 2}}]}
        canon_activity_day(dia)
        self.assertEqual(dia["usuarios"][0]["procs_i"], {"chrome": 5})


class TestCoreSettingsHorario(ObserverTestCase):
    def setUp(self):
        self.authenticate()
        self.setup_coresettings()

    def test_valida_y_normaliza(self):
        url = "/core/settings/"
        r = self.client.put(
            url, {"activity_work_hours": "8:00-17:30", "activity_work_days": "4,0"}
        )
        self.assertEqual(r.status_code, 200, r.data)
        cs = get_core_settings()
        self.assertEqual(cs.activity_work_hours, "08:00-17:30")
        self.assertEqual(cs.activity_work_days, "0,4")
        r = self.client.put(url, {"activity_work_hours": "18:00-08:00"})
        self.assertEqual(r.status_code, 400)


LOCMEM = {"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}}


@override_settings(CACHES=LOCMEM)
class TestMonitoreoEnVivo(ObserverTestCase):
    def setUp(self):
        self.authenticate()
        self.setup_coresettings()
        # La receta fija last_seen al importarse: en una corrida larga el equipo
        # ya saldría «offline» y la captura no se despacharía.
        self.agent = baker.make_recipe(
            "agents.online_agent", last_seen=djangotime.now()
        )
        self.tmp = tempfile.mkdtemp()
        self.override = self.settings(ACTIVITY_SCREENSHOTS_BASE_PATH=self.tmp)
        self.override.enable()
        self.addCleanup(self.override.disable)
        cache.clear()

    def _encender(self, capturas_min=10):
        cs = get_core_settings()
        cs.activity_monitoring_enabled = True
        cs.activity_screenshots_min = capturas_min
        cs.save()
        self.agent.activity_monitoring = True
        self.agent.save()

    def _url(self):
        return f"/agents/{self.agent.agent_id}/activity/screenshots/now/"

    def test_apagado_no_llama_al_agente(self):
        with patch("agents.models.Agent.nats_cmd", new_callable=AsyncMock) as nats:
            r = self.client.post(self._url())
        self.assertEqual(r.data["estado"], "monitoreo_off")
        nats.assert_not_called()
        self._encender(capturas_min=0)
        with patch("agents.models.Agent.nats_cmd", new_callable=AsyncMock) as nats:
            r = self.client.post(self._url())
        self.assertEqual(r.data["estado"], "capturas_off")
        nats.assert_not_called()

    def test_captura_forzada_audita_y_limita(self):
        self._encender()
        with patch(
            "agents.models.Agent.nats_cmd",
            new_callable=AsyncMock,
            return_value={"estado": "bloqueado"},
        ) as nats:
            r = self.client.post(self._url())
            self.assertEqual(r.data["estado"], "bloqueado")
            nats.assert_called_once_with({"func": "actividad_captura"}, timeout=45)
            # Un segundo pedido dentro de 15 s no llega al equipo.
            r = self.client.post(self._url())
            self.assertEqual(r.data["estado"], "espera")
            self.assertEqual(nats.call_count, 1)
        log = AuditLog.objects.filter(agent_id=self.agent.agent_id).last()
        self.assertIn("forced a live screenshot", log.message)
        self.assertIn("bloqueado", log.message)

    def test_sin_conexion_responde_al_tiro(self):
        self._encender()
        self.agent.last_seen = None
        self.agent.save()
        with patch("agents.models.Agent.nats_cmd", new_callable=AsyncMock) as nats:
            r = self.client.post(self._url())
        self.assertEqual(r.data["estado"], "offline")
        nats.assert_not_called()

    def test_sin_permiso_403(self):
        self._encender()
        user = self.create_user_with_roles([])
        self.client.force_authenticate(user=user)
        with patch("agents.models.Agent.nats_cmd", new_callable=AsyncMock) as nats:
            self.check_not_authorized("post", self._url())
            self.check_not_authorized("get", "/agents/activity/live/")
            self.check_not_authorized("get", "/agents/activity/export/")
        nats.assert_not_called()

    def test_vista_en_vivo_lista_equipos_monitoreados(self):
        otro = baker.make_recipe("agents.online_agent")
        self._encender()
        AgentActivityScreenshot.objects.create(
            agent=self.agent,
            captured_at="2026-10-13T10:00:00Z",
            username="ana",
            process="chrome",
            url="docs.google.com/x",
            manual=True,
        )
        r = self.client.get("/agents/activity/live/")
        self.assertEqual(r.status_code, 200)
        ids = [a["agent_id"] for a in r.data]
        self.assertIn(self.agent.agent_id, ids)
        self.assertNotIn(otro.agent_id, ids)
        fila = next(a for a in r.data if a["agent_id"] == self.agent.agent_id)
        self.assertEqual(fila["ultima"]["url"], "docs.google.com/x")


class TestExportXLSX(ObserverTestCase):
    def setUp(self):
        self.authenticate()
        self.setup_coresettings()
        self.agent = baker.make_recipe("agents.online_agent")

    def test_exporta_resumen_apps_y_urls(self):
        from openpyxl import load_workbook

        AgentActivityDaily.objects.create(
            agent=self.agent,
            date="2026-10-13",
            data={
                "v": 1,
                "so": "windows",
                "dias": [
                    {
                        "fecha": "2026-10-13",
                        "usuarios": [
                            {
                                "usuario": "ana",
                                "activo_min": 60,
                                "inactivo_min": 30,
                                "procs": {"chrome": 60},
                                "procs_i": {"chrome": 30},
                                "urls": {"docs.google.com/x": 40},
                                "teclas": 100,
                                "mouse": 50,
                            }
                        ],
                    }
                ],
            },
        )
        r = self.client.get(
            "/agents/activity/export/",
            {
                "client": self.agent.site.client_id,
                "desde": "2026-10-01",
                "hasta": "2026-10-31",
            },
        )
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(
            wb.sheetnames,
            ["Resumen", "Aplicaciones", "Programas y sitios", "URL", "Sitios web"],
        )
        resumen = list(wb["Resumen"].values)
        self.assertEqual(resumen[1][3], "ana")
        self.assertEqual(resumen[1][16:18], (100, 50))
        self.assertEqual(list(wb["URL"].values)[1][0], "docs.google.com/x")
        self.assertEqual(list(wb["Aplicaciones"].values)[1][2:4], (60, 30))

    def test_fechas_invalidas(self):
        r = self.client.get(
            "/agents/activity/export/", {"desde": "2026-10-31", "hasta": "2026-10-01"}
        )
        self.assertEqual(r.status_code, 400)
